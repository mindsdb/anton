"""Coverage for the LibreOffice recalculation oracle.

`check_xlsx_via_office` is the oracle path ahead of the structural lint —
see `xlsx_office_check.py`'s module docstring for why `None` (not `[]`)
means "couldn't run, fall back".
"""

from __future__ import annotations

import os
import shutil
from pathlib import Path

import openpyxl
import pytest

from anton.core.artifacts.xlsx_office_check import check_xlsx_via_office

# Unlike the html lint's browser (bundled inside Electron, unreachable via
# PATH without the env var), a real LibreOffice install is on PATH already —
# so these tests should run whenever soffice is actually available, not only
# when the override env var happens to be set.
requires_office = pytest.mark.skipif(
    not (os.environ.get("ANTON_XLSX_LINT_OFFICE") or shutil.which("soffice") or shutil.which("libreoffice")),
    reason="no soffice/libreoffice binary found on PATH or via ANTON_XLSX_LINT_OFFICE",
)


def _write_workbook(path: Path, sheets: dict[str, dict[str, object]]) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cells in sheets.items():
        ws = wb.create_sheet(title=name)
        for ref, value in cells.items():
            ws[ref] = value
    wb.save(path)


@pytest.fixture
def circular_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "forecast.xlsx"
    _write_workbook(
        path,
        {
            "Actuals": {c: 100 + i * 10 for i, c in enumerate(["E6", "F6", "G6", "H6", "I6", "J6"])},
            # Missing sheet prefix: range is self-referencing, same as the
            # customer's 15 cells in ENG-1204.
            "Assumptions": {"E6": "=SLOPE(E6:J6,{1,2,3,4,5,6})"},
        },
    )
    return path


@pytest.fixture
def clean_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "clean.xlsx"
    _write_workbook(
        path,
        {
            "Actuals": {c: 100 + i * 10 for i, c in enumerate(["E6", "F6", "G6", "H6", "I6", "J6"])},
            "Assumptions": {"E6": "=SLOPE('Actuals'!E6:J6,{1,2,3,4,5,6})"},
        },
    )
    return path


def test_no_office_configured_yields_none(monkeypatch, circular_workbook: Path):
    monkeypatch.delenv("ANTON_XLSX_LINT_OFFICE", raising=False)
    monkeypatch.setattr("shutil.which", lambda _name: None)
    assert check_xlsx_via_office(circular_workbook) is None


def test_office_path_that_does_not_exist_yields_none(monkeypatch, circular_workbook: Path):
    monkeypatch.setenv("ANTON_XLSX_LINT_OFFICE", "/no/such/binary-xyz")
    assert check_xlsx_via_office(circular_workbook) is None


@requires_office
def test_flags_a_formula_that_recalculates_to_an_error(circular_workbook: Path):
    findings = check_xlsx_via_office(circular_workbook)
    assert findings is not None
    assert len(findings) == 1
    assert findings[0].sheet == "Assumptions"
    assert findings[0].cell == "E6"
    assert findings[0].error.startswith("#")
    assert "recalculated" in findings[0].message()


@requires_office
def test_clean_workbook_has_no_findings(clean_workbook: Path):
    assert check_xlsx_via_office(clean_workbook) == []


@pytest.fixture
def string_results_workbook(tmp_path: Path) -> Path:
    """Formulas whose STRING result merely starts with '#'/'Err:' — none of
    these are formula errors, just text a real workbook can legitimately
    produce (a rank badge, a tag, a label someone chose to write)."""
    path = tmp_path / "labels.xlsx"
    _write_workbook(
        path,
        {
            "Sheet1": {
                "B2": 5,
                # Rank badge, e.g. "#3" — a formula concatenating a literal
                # "#" with a computed number.
                "C1": '="#"&TEXT(RANK(B2,B2:B5),"0")',
                "C2": '=CONCATENATE("#tag-",B2)',
                # A literal string that happens to spell an error code.
                "C3": '="Err:"&"522"',
            }
        },
    )
    return path


@requires_office
def test_does_not_flag_a_string_result_that_looks_like_an_error(string_results_workbook: Path):
    """The exact false-positive shape from the review: `_ERROR_PREFIXES`
    string-matching flagged these; only `data_type == 'e'` — a genuine
    formula error, not any string starting with '#' — should."""
    assert check_xlsx_via_office(string_results_workbook) == []
