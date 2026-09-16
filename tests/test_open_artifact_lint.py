"""`open_artifact` re-runs the lint, not only the write-time hook.

Before this, a finding only ever reached the agent in the tool result of
the exact cell that wrote the broken file — persisted nowhere, not re-raised
on a later `open_artifact`. An agent that didn't act on it in that moment
(out of retries, or the artifact deliberately incomplete mid-turn) would
never see it again. This pins the fix: opening an artifact re-checks its
CURRENT files regardless of whether this turn wrote them.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from anton.core.artifacts import ArtifactStore
from anton.core.tools.tool_handlers import handle_open_artifact


@pytest.fixture(autouse=True)
def _no_office(monkeypatch):
    # Deterministic: exercise the structural lint regardless of whether the
    # machine running this happens to have LibreOffice installed.
    monkeypatch.delenv("ANTON_XLSX_LINT_OFFICE", raising=False)
    monkeypatch.setattr("shutil.which", lambda _name: None)


class FakeWorkspace:
    def __init__(self, root: Path) -> None:
        self.artifacts_dir = root


class FakeSession:
    def __init__(self, root: Path) -> None:
        self._workspace = FakeWorkspace(root)
        self._session_id = "conv-1"
        self._turn_count = 0
        self._artifacts_touched: set[str] = set()
        self._data_vault = None


@pytest.fixture
def root(tmp_path: Path) -> Path:
    return tmp_path / "artifacts"


@pytest.fixture
def session(root: Path) -> FakeSession:
    return FakeSession(root)


def _write_broken_workbook(folder: Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    actuals = wb.create_sheet("Actuals")
    actuals["E6"] = 100
    assumptions = wb.create_sheet("Assumptions")
    assumptions["E6"] = "=SLOPE(E6:J6,{1,2,3,4,5,6})"  # missing sheet prefix
    wb.save(folder / "forecast.xlsx")


def _write_clean_workbook(folder: Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    actuals = wb.create_sheet("Actuals")
    for i, col in enumerate("EFGHIJ"):
        actuals[f"{col}6"] = 100 + i * 10
    assumptions = wb.create_sheet("Assumptions")
    assumptions["E6"] = "=SLOPE('Actuals'!E6:J6,{1,2,3,4,5,6})"
    wb.save(folder / "forecast.xlsx")


async def test_reopening_a_broken_artifact_surfaces_the_finding(session, root):
    artifact = ArtifactStore(root).create(name="Forecast", description="d", type="document")
    folder = root / artifact.slug
    _write_broken_workbook(folder)
    # Nothing "just wrote" this from the tool layer's point of view — the
    # write above happened outside any cell this handler observed, same as
    # a file written in an earlier turn/session.

    outcome = await handle_open_artifact(session, {"slug": artifact.slug})

    assert "[artifact lint]" in outcome.content
    assert "E6" in outcome.content
    assert "missing cross-sheet prefix" in outcome.content


async def test_reopening_a_clean_artifact_has_no_lint_section(session, root):
    artifact = ArtifactStore(root).create(name="Forecast", description="d", type="document")
    folder = root / artifact.slug
    _write_clean_workbook(folder)

    outcome = await handle_open_artifact(session, {"slug": artifact.slug})

    assert "[artifact lint]" not in outcome.content
