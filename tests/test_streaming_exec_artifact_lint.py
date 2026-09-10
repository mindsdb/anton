"""The artifact lint hook must fire on the REAL product path.

`ChatSession.turn_stream`'s inline scratchpad exec (session.py) bypasses
`handle_scratchpad` (tool_handlers.py) entirely — confirmed by grep: no
caller reaches `handle_scratchpad`'s exec branch outside tests, since every
real entry point (`anton/chat.py`, `cloud_turn/__main__.py`,
`commands/goal.py`, `dispatch/local_runtime.py`) calls `turn_stream`, never
`turn()`. A lint hook wired only into `handle_scratchpad` would never run in
production. This test pins the inline path itself.

The scratchpad pad is faked (no real subprocess sandbox) so the test stays
fast and doesn't depend on a venv having `openpyxl`/LibreOffice installed —
what's under test is "does turn_stream call the lint hook and surface its
findings", not the checkers themselves (already covered by
test_xlsx_lint.py / test_tool_handlers_xlsx_lint.py).
"""

from __future__ import annotations

from pathlib import Path
from unittest.mock import AsyncMock, MagicMock

import openpyxl
import pytest
from tests.conftest import make_mock_llm

from anton.core.artifacts import ArtifactStore
from anton.core.backends.base import Cell
from anton.core.llm.provider import LLMResponse, StreamComplete, ToolCall, Usage
from anton.core.session import ChatSession, ChatSessionConfig, _VerifierVerdict


@pytest.fixture()
def workspace(tmp_path: Path):
    ws = MagicMock(base=tmp_path)
    ws.artifacts_dir = tmp_path / "artifacts"
    return ws


class _FakeAsyncIter:
    def __init__(self, items):
        self._items = items

    def __aiter__(self):
        return self

    async def __anext__(self):
        if not self._items:
            raise StopAsyncIteration
        return self._items.pop(0)


class _FakePad:
    """Stands in for a real ScratchpadRuntime: no subprocess, no venv —
    `execute_streaming` just writes the broken workbook straight to disk
    (the side effect the real sandbox would have produced) and yields the
    Cell the real runtime would return."""

    def __init__(self, workbook_path: Path) -> None:
        self._workbook_path = workbook_path

    async def install_packages(self, packages):  # pragma: no cover - unused here
        return "nothing to install"

    async def execute_streaming(self, code, **kwargs):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        actuals = wb.create_sheet("Actuals")
        actuals["E6"] = 100
        assumptions = wb.create_sheet("Assumptions")
        assumptions["E6"] = "=SLOPE(E6:J6,{1,2,3,4,5,6})"  # missing sheet prefix
        wb.save(self._workbook_path)
        yield Cell(
            code=code, stdout="saved forecast.xlsx", stderr="", error=None,
            description="build forecast",
        )


def _text_response(text: str) -> LLMResponse:
    return LLMResponse(
        content=text, tool_calls=[], usage=Usage(input_tokens=10, output_tokens=20),
        stop_reason="end_turn",
    )


def _scratchpad_exec_response(text: str) -> LLMResponse:
    return LLMResponse(
        content=text,
        tool_calls=[
            ToolCall(
                id="tc_sp_1", name="scratchpad",
                input={"action": "exec", "name": "main", "code": "build_forecast()"},
            ),
        ],
        usage=Usage(input_tokens=10, output_tokens=20),
        stop_reason="tool_use",
    )


async def test_turn_stream_surfaces_a_lint_finding_from_inline_exec(workspace):
    artifact = ArtifactStore(workspace.artifacts_dir).create(
        name="Forecast", description="d", type="document"
    )
    workbook_path = workspace.artifacts_dir / artifact.slug / "forecast.xlsx"

    mock_llm = make_mock_llm()
    mock_llm.generate_object_code = AsyncMock(
        return_value=_VerifierVerdict(status="COMPLETE", reason="task done")
    )
    call_count = 0

    def fake_plan_stream(**kwargs):
        nonlocal call_count
        call_count += 1
        if call_count == 1:
            return _FakeAsyncIter(
                [StreamComplete(response=_scratchpad_exec_response("Building the workbook."))]
            )
        return _FakeAsyncIter([StreamComplete(response=_text_response("Done."))])

    mock_llm.plan_stream = fake_plan_stream

    session = ChatSession(ChatSessionConfig(llm_client=mock_llm, workspace=workspace))
    session._scratchpads.get_or_create = AsyncMock(return_value=_FakePad(workbook_path))
    try:
        async for _ in session.turn_stream("build the forecast workbook"):
            pass

        tool_result_msgs = [
            m for m in session.history
            if m["role"] == "user" and isinstance(m["content"], list)
        ]
        assert tool_result_msgs, "the exec call must produce a tool_result"
        result_content = tool_result_msgs[0]["content"][0]["content"]
        # In-memory, per-turn status a host reads to decide whether to warn
        # the user at end of turn — not persisted anywhere past the turn.
        lint_status = session.artifact_lint_status
    finally:
        await session.close()

    assert "[artifact lint]" in result_content
    assert "E6" in result_content
    assert lint_status == {artifact.slug: "has_errors"}
