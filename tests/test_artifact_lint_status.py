"""Per-artifact lint status for the end-of-turn user notice.

`lint_changed_artifact_files(..., status_by_slug=...)` collapses per-file
checker messages into one of two states per artifact — `has_errors` (a real
finding) outranks `not_validated` (the checker never actually ran) — or
clears the slug's entry when a re-lint comes back clean. In-memory, per
turn (see `ChatSession.artifact_lint_status`), no disk write.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from anton.core.tools.tool_handlers import (
    LINT_STATUS_HAS_ERRORS,
    LINT_STATUS_NOT_VALIDATED,
    lint_changed_artifact_files,
)


@pytest.fixture(autouse=True)
def _no_office(monkeypatch):
    monkeypatch.delenv("ANTON_XLSX_LINT_OFFICE", raising=False)
    monkeypatch.setattr("shutil.which", lambda _name: None)


def _write_workbook(path: Path, sheets: dict[str, dict[str, str]]) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, cells in sheets.items():
        ws = wb.create_sheet(title=name)
        for ref, formula in cells.items():
            ws[ref] = f"={formula}"
    wb.save(path)


class _FakeStore:
    def __init__(self, root: Path) -> None:
        self.root = root


@pytest.fixture
def store(tmp_path: Path) -> _FakeStore:
    return _FakeStore(tmp_path)


def _make_artifact(store: _FakeStore, slug: str) -> Path:
    folder = store.root / slug
    folder.mkdir()
    (folder / "metadata.json").write_text("{}")
    return folder


def test_has_errors_when_a_real_finding_is_present(store: _FakeStore):
    folder = _make_artifact(store, "forecast-abc12345")
    _write_workbook(
        folder / "forecast.xlsx",
        {
            "Actuals": {"E6": "100"},
            "Assumptions": {"E6": "SLOPE(E6:J6,{1,2,3,4,5,6})"},  # missing sheet prefix
        },
    )
    status: dict[str, str] = {}

    lint_changed_artifact_files(store, {"forecast-abc12345": 0.0}, status_by_slug=status)

    assert status == {"forecast-abc12345": LINT_STATUS_HAS_ERRORS}


def test_no_browser_at_all_sets_no_status(monkeypatch, store: _FakeStore):
    """No browser configured for this deployment at all (cloud/web) is a
    capability fact, not a per-artifact finding — `.html` isn't registered
    as a checker here, so it never contributes `not_validated` either.
    Otherwise every html artifact on such a host would carry a permanent,
    uncloseable 'Not validated' badge."""
    monkeypatch.delenv("ANTON_HTML_LINT_BROWSER", raising=False)
    folder = _make_artifact(store, "dash-abc12345")
    (folder / "dash.html").write_text("<html></html>")
    status: dict[str, str] = {}

    lint_changed_artifact_files(store, {"dash-abc12345": 0.0}, status_by_slug=status)

    assert status == {}


def test_not_validated_when_a_configured_checker_could_not_run(monkeypatch, store: _FakeStore):
    """`.xlsx` always has the structural lint as a fallback (it never
    returns None). For `.html`, this state now arises only when a browser
    IS configured but this particular run didn't produce a result (crash/
    timeout/malformed output) — distinct from no browser being configured
    at all, which sets no status (see the test above)."""
    monkeypatch.setenv("ANTON_HTML_LINT_BROWSER", "/bin/true")  # exits 0, no output
    folder = _make_artifact(store, "dash-abc12345")
    (folder / "dash.html").write_text("<html></html>")
    status: dict[str, str] = {}

    lint_changed_artifact_files(store, {"dash-abc12345": 0.0}, status_by_slug=status)

    assert status == {"dash-abc12345": LINT_STATUS_NOT_VALIDATED}


def test_status_clears_once_the_artifact_re_lints_clean(monkeypatch, store: _FakeStore):
    """The scenario this exists for: the agent fixes the file later in the
    same turn — a later clean re-lint must remove the earlier verdict, not
    leave it stuck at 'has_errors' for the rest of the turn."""
    folder = _make_artifact(store, "forecast-abc12345")
    workbook = folder / "forecast.xlsx"
    _write_workbook(
        workbook,
        {
            "Actuals": {"E6": "100"},
            "Assumptions": {"E6": "SLOPE(E6:J6,{1,2,3,4,5,6})"},
        },
    )
    status: dict[str, str] = {}
    lint_changed_artifact_files(store, {"forecast-abc12345": 0.0}, status_by_slug=status)
    assert status == {"forecast-abc12345": LINT_STATUS_HAS_ERRORS}

    # Office becomes available and reports the (now-fixed) file as clean —
    # simulates the agent fixing the formula and the oracle confirming it.
    monkeypatch.setattr(
        "anton.core.artifacts.xlsx_office_check.check_xlsx_via_office", lambda _path: []
    )
    _write_workbook(workbook, {"Assumptions": {"E6": "1"}})
    before = {"forecast-abc12345": workbook.stat().st_mtime - 1}

    lint_changed_artifact_files(store, before, status_by_slug=status)

    assert status == {}


def test_has_errors_outranks_not_validated_across_sibling_files(monkeypatch, store: _FakeStore):
    """One sibling file only got as far as "could not run" (a configured
    browser whose run produced no result for its `.html`), another has a
    real `.xlsx` finding — the artifact-level status must reflect the worse
    one, regardless of file processing order."""
    monkeypatch.setenv("ANTON_HTML_LINT_BROWSER", "/bin/true")  # exits 0, no output
    folder = _make_artifact(store, "mixed-abc12345")
    _write_workbook(
        folder / "broken.xlsx",
        {
            "Actuals": {"E6": "100"},
            "Assumptions": {"E6": "SLOPE(E6:J6,{1,2,3,4,5,6})"},
        },
    )
    (folder / "dash.html").write_text("<html></html>")
    status: dict[str, str] = {}

    lint_changed_artifact_files(store, {"mixed-abc12345": 0.0}, status_by_slug=status)

    assert status == {"mixed-abc12345": LINT_STATUS_HAS_ERRORS}


def test_status_by_slug_defaults_to_none_and_is_optional(store: _FakeStore):
    """Existing callers that don't pass status_by_slug must keep working
    unchanged — this is an additive, backward-compatible parameter."""
    folder = _make_artifact(store, "forecast-abc12345")
    _write_workbook(
        folder / "forecast.xlsx",
        {
            "Actuals": {"E6": "100"},
            "Assumptions": {"E6": "SLOPE(E6:J6,{1,2,3,4,5,6})"},
        },
    )

    messages = lint_changed_artifact_files(store, {"forecast-abc12345": 0.0})

    assert any("E6" in m for m in messages)
